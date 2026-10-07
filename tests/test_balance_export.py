"""The workbook handed back to the caller.

The first sheet's header is the contract - two rows, with ``จำหน่ายให้`` merged
over its two sub-columns and ``จำนวน`` over its four, exactly as
``ฟอแมทตาราง OCR.xlsx`` sets them out. It is asserted literally rather than
derived from the module under test, so changing the agreed shape takes a
deliberate edit in both places.
"""
import datetime as dt
import io

from django.test import SimpleTestCase
from openpyxl import load_workbook

from workers import balance_export, form_layout
from workers.balance_parser import LedgerRow
from workers.balance_pipeline import PageOutcome

#: The first header row as a reader gets it back: a merged span keeps its text
#: in its top-left cell and reports None for the rest, which is exactly how
#: ``ฟอแมทตาราง OCR.xlsx`` itself reads.
EXPECTED_TOP = [
    "ลำดับ",
    "วัน\nเดือน\nปี",
    "ชื่อ/ความแรงของวัตถุออกฤทธิ์",
    "ชื่อการค้า",
    "เลขที่/รุ่นที่/\nครั้งที่ผลิต",
    "ได้มาจาก",
    "จำหน่ายให้",
    None,
    "จำนวน",
    None,
    None,
    None,
]

EXPECTED_SUB = [
    None, None, None, None, None, None,
    "ชื่อ-นามสกุล\nผู้รับยา",
    "เลขที่บัตรประจำตัวประชาชน/\nหนังสือเดินทาง/บัตรประจำตัวอื่น\n"
    "ที่ทางราชการออกให้",
    "ยกมา",
    "รับ",
    "จ่าย",
    "คงเหลือ",
]

#: How the two header rows are merged: each single-heading column down over
#: both rows, each shared heading across its sub-columns.
EXPECTED_MERGES = {
    "A1:A2", "B1:B2", "C1:C2", "D1:D2", "E1:E2", "F1:F2", "G1:H1", "I1:L1",
}



def _row():
    row = LedgerRow(
        page_number=4,
        row_index=0,
        form_code="บ.ย.ส. ๒/ว.จ. ๒-จ๑",
        cells={
            form_layout.DATE: "5 มค 68",
            form_layout.GENERIC_NAME: "Methylphenidate HCl tablets 10 mg",
            form_layout.TRADE_NAME: "Ritalin tablets 10 mg",
            form_layout.BATCH_NO: "BE210",
            form_layout.RECEIVED_FROM: "อย",
            form_layout.ISSUED_TO: "ด.ญ.พิมพ์ เก่งการดี",
            form_layout.RECIPIENT_ID: "1112223334440",
        },
        entry_date=dt.date(2025, 1, 5),
        confidence=0.88,
    )
    row.quantities = {
        form_layout.BALANCE_BROUGHT: 450,
        form_layout.RECEIVED: None,
        form_layout.ISSUED: 40,
        form_layout.BALANCE: 410,
    }
    row.flag("issued_to is written in Thai, which this engine cannot read")
    return row


def _workbook(rows=None, pages=None):
    payload = balance_export.build(
        rows if rows is not None else [_row()],
        pages
        if pages is not None
        else [PageOutcome(page_number=4, form_code="บ.ย.ส. ๒/ว.จ. ๒-จ๑", rows=1)],
    )
    return load_workbook(io.BytesIO(payload))


def _values(sheet, row=3):
    """One data row. The data starts at row 3: the header takes two."""
    return list(next(sheet.iter_rows(min_row=row, max_row=row, values_only=True)))


class OutputSheetTests(SimpleTestCase):
    def test_header_matches_the_agreed_columns_exactly(self):
        sheet = _workbook()["Output"]
        top = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        sub = [c.value for c in next(sheet.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(top, EXPECTED_TOP)
        # A merged cell reads back as None below its top-left corner, so the
        # second row carries only the sub-headings.
        self.assertEqual([value or None for value in sub], EXPECTED_SUB)

    def test_the_header_spans_are_merged_as_the_format_sets_them(self):
        sheet = _workbook()["Output"]
        self.assertEqual({str(span) for span in sheet.merged_cells.ranges}, EXPECTED_MERGES)

    def test_one_row_per_entry_below_the_two_header_rows(self):
        self.assertEqual(_workbook()["Output"].max_row, 3)

    def test_carries_what_the_form_says(self):
        values = _values(_workbook()["Output"])
        self.assertEqual(values[0], 1)                              # ลำดับ
        self.assertEqual(values[1], "5 มค 68")                      # วัน เดือน ปี
        self.assertEqual(values[2], "Methylphenidate HCl tablets 10 mg")
        self.assertEqual(values[3], "Ritalin tablets 10 mg")
        self.assertEqual(values[4], "BE210")
        self.assertEqual(values[5], "อย")
        self.assertEqual(values[6], "ด.ญ.พิมพ์ เก่งการดี")
        self.assertEqual(values[7], "1112223334440")

    def test_carries_the_four_quantities(self):
        values = _values(_workbook()["Output"])
        self.assertEqual(values[8], 450)    # ยกมา
        self.assertIsNone(values[9])        # รับ - a dash on the form
        self.assertEqual(values[10], 40)    # จ่าย
        self.assertEqual(values[11], 410)   # คงเหลือ

    def test_the_date_is_the_cell_not_a_date_value(self):
        # "3-31 ม.ค. 68" is a month of dispensing on one line, which no date can
        # hold; the column carries what the form writes.
        row = _row()
        row.cells[form_layout.DATE] = "3-31 ม.ค. 68"
        row.entry_date = None
        self.assertEqual(_values(_workbook(rows=[row])["Output"])[1], "3-31 ม.ค. 68")

    def test_the_sequence_restarts_at_each_new_form(self):
        first, second = _row(), _row()
        second.form_code = "ร.ย.ส. ๒/ว.จ. ๒-จ๑"
        third = _row()
        third.form_code = "ร.ย.ส. ๒/ว.จ. ๒-จ๑"
        sheet = _workbook(rows=[first, second, third])["Output"]
        self.assertEqual([_values(sheet, r)[0] for r in (3, 4, 5)], [1, 1, 2])


class WorkbookShapeTests(SimpleTestCase):
    """Two sheets: the agreed columns, and how each page went."""

    def test_there_is_no_sheet_of_raw_detail(self):
        self.assertEqual(_workbook().sheetnames, ["Output", "สรุปรายหน้า"])

    def test_a_row_needing_review_is_shaded_rather_than_explained(self):
        # With the detail sheet gone this is all that is left of the verdict in
        # the workbook, and it is enough to send a reader to the right line.
        sheet = _workbook()["Output"]
        filled = {cell.fill.fgColor.rgb for cell in sheet[3]}
        self.assertEqual(len(filled), 1)
        plain = _row()
        plain.needs_review = False
        plain.review_notes = []
        unshaded = _workbook(rows=[plain])["Output"]
        self.assertNotEqual(
            next(iter(filled)), unshaded[3][0].fill.fgColor.rgb
        )


class PageSheetTests(SimpleTestCase):
    def test_lists_every_page_including_the_unreadable_ones(self):
        pages = [
            PageOutcome(page_number=1, form_code="ร.ค.-๔", rows=6, confidence=0.78),
            PageOutcome(page_number=2, error="no table grid: found 0 columns"),
        ]
        sheet = _workbook(pages=pages)["สรุปรายหน้า"]
        self.assertEqual(sheet.max_row, 3)
        self.assertIn("no table grid", sheet.cell(row=3, column=6).value)

    def test_accepts_the_stored_page_summary_as_plain_dicts(self):
        # The export runs from BalanceDocument.page_summary, which is JSON.
        sheet = _workbook(pages=[{"page": 1, "form": "ร.ค.-๔", "rows": 6}])["สรุปรายหน้า"]
        self.assertEqual(sheet.cell(row=2, column=2).value, "ร.ค.-๔")


class EmptyDocumentTests(SimpleTestCase):
    def test_a_document_with_no_rows_still_produces_the_header(self):
        sheet = _workbook(rows=[], pages=[])["Output"]
        top = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        self.assertEqual(top, EXPECTED_TOP)
        self.assertEqual(sheet.max_row, 2)
