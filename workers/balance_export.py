"""Excel output for a psychotropic ledger extraction.

Sheet 1 carries the twelve columns of ``ฟอแมทตาราง OCR.xlsx`` exactly: the same
headings, the same order, and the same two-row header with ``จำหน่ายให้`` spanning
its two sub-columns and ``จำนวน`` its four.

Unlike the sixteen-column shape this replaced - which described a
purchase-approval record and had somewhere to put only four of the things a
ledger says - these columns *are* the ledger, so eleven of the twelve come
straight off the form. Only ``ลำดับ`` is this service's own: the forms do not
number their lines, so it is counted here, restarting at each new form.

Sheet 2 is the per-page outcome, where a page that could not be read is
visible. There is no third sheet of raw detail any more: the agreed twelve
columns are the deliverable, and a row the extraction is unsure of is shaded
rather than explained. The unit, the remark and the review notes are still read
and still stored - ``apps/balance/models.py`` - they are simply not written out.
"""
from __future__ import annotations

import io
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from workers import form_layout

#: Sheet 1, verbatim from ฟอแมทตาราง OCR.xlsx: (top heading, sub-heading, width).
#: A sub-heading of None means the column has a single heading spanning both
#: header rows. The line breaks are the customer's and are kept, because the
#: JSON keys are these strings.
OUTPUT_COLUMNS: Sequence[tuple[str, "str | None", int]] = (
    ("ลำดับ", None, 7),
    ("วัน\nเดือน\nปี", None, 14),
    ("ชื่อ/ความแรงของวัตถุออกฤทธิ์", None, 30),
    ("ชื่อการค้า", None, 22),
    ("เลขที่/รุ่นที่/\nครั้งที่ผลิต", None, 14),
    ("ได้มาจาก", None, 14),
    ("จำหน่ายให้", "ชื่อ-นามสกุล\nผู้รับยา", 24),
    (
        "จำหน่ายให้",
        "เลขที่บัตรประจำตัวประชาชน/\nหนังสือเดินทาง/บัตรประจำตัวอื่น\n"
        "ที่ทางราชการออกให้",
        30,
    ),
    ("จำนวน", "ยกมา", 10),
    ("จำนวน", "รับ", 10),
    ("จำนวน", "จ่าย", 10),
    ("จำนวน", "คงเหลือ", 10),
)

#: The key each column is given in the API's JSON: the sub-heading where there
#: is one, the heading otherwise, so every key is distinct.
OUTPUT_KEYS: Sequence[str] = tuple(
    sub if sub else top for top, sub, _ in OUTPUT_COLUMNS
)

PAGE_COLUMNS: Sequence[tuple[str, int]] = (
    ("หน้า", 8),
    ("แบบฟอร์ม", 16),
    ("จำนวนแถว", 11),
    ("เอียง (องศา)", 13),
    ("Confidence", 12),
    ("ปัญหา", 60),
)

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")


def build(rows, pages) -> bytes:
    """Return the workbook for one extraction.

    ``rows`` are ``balance_parser.LedgerRow`` or ``BalanceEntry`` instances -
    anything exposing the same attribute names - and ``pages`` are
    ``balance_pipeline.PageOutcome``.
    """
    workbook = Workbook()

    output = workbook.active
    output.title = "Output"
    _write_output(output, rows)

    _write_pages(workbook.create_sheet("สรุปรายหน้า"), pages)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def output_values(row, sequence: int) -> list:
    """The twelve agreed columns for one entry, in order.

    The single definition of that mapping. The workbook's first sheet and the
    API's JSON are both built from it, so the two cannot drift apart - which
    they would, being written in different modules against the same agreed list.

    ``sequence`` is the ลำดับ, which the forms do not carry; see ``sequences``.

    The date is the cell as the form writes it, not a date value: these forms
    use Thai month abbreviations and, on the report form, ranges - "3-31 ม.ค.
    68" covers a whole month of dispensing on one line. A range is not a date
    and squeezing it into one would lose what the line actually says.
    """
    return [
        sequence,
        _value(row, "entry_date_text") or _text(row, form_layout.DATE),
        _text(row, form_layout.GENERIC_NAME),
        _text(row, form_layout.TRADE_NAME),
        _text(row, form_layout.BATCH_NO),
        _text(row, form_layout.RECEIVED_FROM),
        _text(row, form_layout.ISSUED_TO),
        _text(row, form_layout.RECIPIENT_ID),
        _quantity(row, form_layout.BALANCE_BROUGHT),
        _quantity(row, form_layout.RECEIVED),
        _quantity(row, form_layout.ISSUED),
        _quantity(row, form_layout.BALANCE),
    ]


def sequences(rows) -> list[int]:
    """The ลำดับ of each row: a running count that restarts at each new form.

    The forms do not number their own lines, and a submission holds the returns
    of more than one of them, so one number running through the file would say
    nothing. Restarting per form makes it the line number within that form,
    which is what a reader checking against the paper needs.
    """
    numbers: list[int] = []
    current: str | None = None
    count = 0
    for row in rows:
        form = _value(row, "form_code")
        if form != current:
            current, count = form, 0
        count += 1
        numbers.append(count)
    return numbers


def output_row(row, sequence: int) -> dict:
    """The twelve agreed columns for one entry, keyed by their heading."""
    return dict(zip(OUTPUT_KEYS, output_values(row, sequence)))


def _write_output(sheet: Worksheet, rows) -> None:
    _header(sheet, OUTPUT_COLUMNS)
    for row, sequence in zip(rows, sequences(rows)):
        sheet.append(output_values(row, sequence))
        if _value(row, "needs_review"):
            for cell in sheet[sheet.max_row]:
                cell.fill = _REVIEW_FILL

    _finish(sheet, OUTPUT_COLUMNS, header_rows=2)


def _write_pages(sheet: Worksheet, pages) -> None:
    _header(sheet, PAGE_COLUMNS)
    for page in pages:
        data = page.as_dict() if hasattr(page, "as_dict") else dict(page)
        sheet.append(
            [
                data.get("page"),
                data.get("form"),
                data.get("rows"),
                data.get("skew"),
                data.get("confidence"),
                data.get("error"),
            ]
        )
        if data.get("error"):
            for cell in sheet[sheet.max_row]:
                cell.fill = _REVIEW_FILL
    _finish(sheet, PAGE_COLUMNS)


# --- value access ----------------------------------------------------------
#
# The same workbook is built from the parser's LedgerRow (cells keyed by role)
# and from the stored BalanceEntry (one attribute per column), so every read
# goes through these.


def _value(row, name: str):
    if hasattr(row, name):
        return getattr(row, name) or ("" if name != "needs_review" else False)
    cells = getattr(row, "cells", None)
    return (cells or {}).get(name, "")


def _text(row, role: str) -> str:
    """A cell's text, whether the row is a parsed one or a stored one.

    The parser keys its cells by column role and the stored row has one
    attribute per column; the roles are named after those attributes, so the
    same name reaches both.
    """
    return _value(row, role)


def _quantity(row, role: str):
    quantities = getattr(row, "quantities", None)
    if quantities is not None:
        return quantities.get(role)
    return getattr(row, role, None)




# --- sheet furniture -------------------------------------------------------


def _header(sheet: Worksheet, columns: Sequence[tuple]) -> None:
    """Write a header of one row, or of two with the spans merged.

    A three-part column - (heading, sub-heading, width) - makes a two-row
    header: a column with no sub-heading is merged down over both rows, and
    neighbours sharing a heading are merged across, which is how
    ``ฟอแมทตาราง OCR.xlsx`` sets out จำหน่ายให้ and จำนวน.
    """
    two_row = len(columns[0]) == 3
    if not two_row:
        sheet.append([title for title, _ in columns])
    else:
        sheet.append([top if sub is None else top for top, sub, _ in columns])
        sheet.append(["" if sub is None else sub for _, sub, _ in columns])

    rows = 2 if two_row else 1
    for index, column in enumerate(columns, start=1):
        width = column[-1]
        sheet.column_dimensions[get_column_letter(index)].width = width
        for row in range(1, rows + 1):
            cell = sheet.cell(row=row, column=index)
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
    for row in range(1, rows + 1):
        sheet.row_dimensions[row].height = 30 if not two_row else 26

    if two_row:
        _merge_header(sheet, columns)


def _merge_header(sheet: Worksheet, columns: Sequence[tuple]) -> None:
    index = 1
    while index <= len(columns):
        top, sub, _ = columns[index - 1]
        if sub is None:
            sheet.merge_cells(start_row=1, start_column=index, end_row=2, end_column=index)
            index += 1
            continue
        span = index
        while span < len(columns) and columns[span][0] == top:
            span += 1
        sheet.merge_cells(start_row=1, start_column=index, end_row=1, end_column=span)
        index = span + 1


def _finish(
    sheet: Worksheet, columns: Sequence[tuple], header_rows: int = 1
) -> None:
    sheet.freeze_panes = f"A{header_rows + 1}"
    last = get_column_letter(len(columns))
    # The filter goes on the row the data is actually labelled by: with a merged
    # two-row header that is the second row, not the first.
    sheet.auto_filter.ref = (
        f"A{header_rows}:{last}{max(sheet.max_row, header_rows)}"
    )
