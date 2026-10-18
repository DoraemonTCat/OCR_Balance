"""Write the answer key and the key-vs-OCR comparison as Excel workbooks.

    python eval/report.py eval/out/final.json eval/เฉลย_ต.ย.OCR.xlsx eval/ผลเทียบ_OCR_กับเฉลย.xlsx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from eval.compare import FIELDS, align, field_equal, load_key  # noqa: E402

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True)
BAD = PatternFill("solid", fgColor="F8CBAD")
REVIEW = PatternFill("solid", fgColor="FFF2CC")
GOOD = PatternFill("solid", fgColor="E2EFDA")

LABELS = {
    "date": "วัน เดือน ปี", "name": "ชื่อผู้รับยา", "id": "เลขบัตรประชาชน",
    "brought": "ยอดยกมา", "received": "รับ", "issued": "จ่าย", "balance": "คงเหลือ",
    "trade_name": "ชื่อการค้า",
}


def _header(ws, titles, widths):
    ws.append(titles)
    for cell, width in zip(ws[1], widths):
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def write_key(pages, out: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "เฉลย"
    _header(
        ws,
        ["หน้า", "แบบ", "ชื่อการค้า", "เดือน", "ปี (หัวกระดาษ)", "ลำดับ", "วัน เดือน ปี (ตามที่พิมพ์)",
         "วัน เดือน ปี (มีผล)", "ชื่อผู้รับยา", "เลขบัตรประชาชน", "อายุ", "ยอดยกมา", "รับ", "จ่าย",
         "คงเหลือ", "หมายเหตุ"],
        [6, 13, 26, 11, 9, 7, 20, 20, 34, 20, 6, 9, 8, 8, 9, 40],
    )
    totals = wb.create_sheet("ยอดรวมท้ายหน้า")
    _header(totals, ["หน้า", "ยอดยกมา", "รับ", "จ่าย", "คงเหลือ", "ผู้รับอนุญาต"], [6, 10, 10, 10, 10, 20])
    for page in pages:
        key = load_key(page)
        for index, row in enumerate(key["rows"], 1):
            ws.append([
                page, key["form_code"], key["trade_name"], key["month"], key["year"], index,
                row["date"], row["effective_date"], row["name"], row["id"], row["age"],
                row["brought"], row["received"], row["issued"], row["balance"], row["note"],
            ])
        t = key["totals"]
        totals.append([page, t["brought"], t["received"], t["issued"], t["balance"], t["signer"]])
    wb.save(out)


def write_comparison(result, pages, out: Path):
    wb = Workbook()
    summary = wb.active
    summary.title = "สรุป"
    rows_ws = wb.create_sheet("รายแถว")
    _header(
        rows_ws,
        ["หน้า", "แถว", "ผล"] + [f"{LABELS[f]} (เฉลย)" for f in FIELDS[:3]]
        + [f"{LABELS[f]} (OCR)" for f in FIELDS[:3]]
        + ["ตัวเลข ยกมา/รับ/จ่าย/คงเหลือ (เฉลย)", "ตัวเลข (OCR)", "ชื่อการค้า (OCR)", "ต้องตรวจทาน", "เหตุผล"],
        [6, 6, 10, 20, 32, 20, 20, 32, 20, 24, 24, 26, 10, 50],
    )
    correct = {f: 0 for f in FIELDS}
    total = perfect = review = wrong_unflagged = 0
    for page in pages:
        key = load_key(page)
        got = [r for r in result["rows"] if r["page"] == page]
        for index, (want, have) in enumerate(align(key["rows"], got), 1):
            total += 1
            if have is None:
                rows_ws.append([page, index, "ขาดหาย", want["effective_date"], want["name"], want["id"]])
                continue
            if want is None:
                rows_ws.append([page, index, "เกินมา", "", "", "", have["date"], have["name"], have["id"]])
                continue
            ok = {}
            for f in FIELDS:
                w = want["effective_date"] if f == "date" else want.get(f)
                ok[f] = field_equal(f, w, have.get(f))
                correct[f] += ok[f]
            all_ok = all(ok.values())
            perfect += all_ok
            review += bool(have["needs_review"])
            wrong_unflagged += (not all_ok) and not have["needs_review"]
            qty = lambda r: "/".join("-" if r[k] is None else str(r[k]) for k in ("brought", "received", "issued", "balance"))
            rows_ws.append([
                page, index, "ถูกทุกช่อง" if all_ok else "มีช่องผิด",
                want["effective_date"], want["name"], want["id"],
                have["date"], have["name"], have["id"],
                qty(want), qty(have), have["trade_name"],
                "ใช่" if have["needs_review"] else "", "; ".join(have["notes"] or []),
            ])
            line = rows_ws.max_row
            rows_ws.cell(line, 3).fill = GOOD if all_ok else BAD
            for col, f in ((7, "date"), (8, "name"), (9, "id"), (12, "trade_name")):
                if not ok[f]:
                    rows_ws.cell(line, col).fill = BAD
            if not all(ok[f] for f in ("brought", "received", "issued", "balance")):
                rows_ws.cell(line, 11).fill = BAD
            if have["needs_review"]:
                rows_ws.cell(line, 13).fill = REVIEW

    _header(summary, ["ช่อง", "ถูก", "ทั้งหมด", "ร้อยละ"], [26, 10, 10, 10])
    for f in FIELDS:
        summary.append([LABELS[f], correct[f], total, round(100 * correct[f] / total, 2)])
    summary.append([])
    summary.append(["แถวที่ถูกทุกช่อง", perfect, total, round(100 * perfect / total, 2)])
    summary.append(["แถวที่ระบบติดธงให้ตรวจทาน", review, total, round(100 * review / total, 2)])
    summary.append(["แถวที่ผิดแต่ไม่ถูกติดธง", wrong_unflagged, total, round(100 * wrong_unflagged / total, 2)])
    summary.append(["เวลาประมวลผลทั้งไฟล์ (วินาที)", round(result["processing_ms"] / 1000, 1)])
    wb.save(out)
    return correct, total, perfect, review, wrong_unflagged


if __name__ == "__main__":
    result = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    pages = sorted({p["page"] for p in result["pages"]})
    write_key(pages, Path(sys.argv[2]))
    print(write_comparison(result, pages, Path(sys.argv[3])))
